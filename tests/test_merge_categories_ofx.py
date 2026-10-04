import calendar
import re
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from stmtconv.core.categorize import Rule, categorize
from stmtconv.core.merge import merge
from stmtconv.core.models import Statement, StatementSummary, Transaction
from stmtconv.export import merged, ofx


def monthly(month):
    summary = StatementSummary(
        account_mask="****1234",
        account_key="a" * 64,
        currency="USD",
        direction="asset",
        period_start=date(2026, month, 1),
        period_end=date(2026, month, calendar.monthrange(2026, month)[1]),
        opening=Decimal(1000 + month - 1),
        closing=Decimal(1000 + month),
        total_debits=Decimal(0),
        total_credits=Decimal(1),
    )
    return Statement(
        id=f"s{month:04d}",
        summary=summary,
        transactions=[
            Transaction(
                id=f"row-{month}",
                date=date(2026, month, 1),
                description="Fake deposit",
                credit=Decimal(1),
                balance=summary.closing,
                source_file="fake.pdf",
                page=1,
                engine="text",
                raw_text="fake",
            )
        ],
        verdict="VERIFIED",
    )


def test_twelve_months_workbook_and_continuity(tmp_path):
    statements = [monthly(m) for m in range(1, 13)]
    result = merge(statements, Decimal("0.01"))
    assert result.issues == []
    path = tmp_path / "merged.xlsx"
    merged.write(result, path)
    book = load_workbook(path)
    assert all(f"2026-{m:02d}" in book.sheetnames for m in range(1, 13))
    assert book["All Transactions"].max_row == 13
    statements[6].summary.opening += Decimal("12.34")
    issues = merge([s for s in statements if s.id != "s0005"], Decimal("0.01")).issues
    assert any(i.code == "PERIOD_GAP" and i.statement_id == "s0006" for i in issues)
    assert any(
        i.code == "CONTINUITY_GAP"
        and i.statement_id == "s0007"
        and i.difference == Decimal("12.34")
        for i in issues
    )


def test_merge_rejects_unknown_or_different_account():
    first, second = monthly(1), monthly(2)
    second.summary.account_mask = "****9999"
    assert any(
        i.code == "ACCOUNT_GROUP_UNKNOWN" for i in merge([first, second], Decimal("0.01")).issues
    )


def test_category_order_direction_and_override():
    statement = monthly(1)
    rules = [
        Rule("wrong direction", ["deposit"], "debit"),
        Rule("First", ["deposit"]),
        Rule("Second", ["deposit"]),
    ]
    assert categorize(statement, rules).transactions[0].category == "First"
    assert (
        categorize(statement, [Rule("Client override", ["deposit"]), *rules])
        .transactions[0]
        .category
        == "Client override"
    )
    assert (
        categorize(statement, [Rule("Regex", ["regex:^Fake"])]).transactions[0].category == "Regex"
    )
    assert statement.transactions[0].category is None


@pytest.mark.parametrize("direction", ["asset", "liability"])
def test_ofx_is_deterministic_escaped_unique_and_roundtrips(direction):
    statement = monthly(1)
    statement.summary.direction = direction
    statement.transactions[0].description = "Fake <merchant> & café"
    duplicate = statement.transactions[0].model_copy(update={"id": "duplicate"}, deep=True)
    statement.transactions.append(duplicate)
    first = ofx.serialize(statement, "20260101000000")
    assert first == ofx.serialize(statement, "20260101000000")
    text = first.decode("cp1252")
    ids = re.findall(r"<FITID>([^<\r\n]+)", text)
    assert len(ids) == len(set(ids)) == 2
    amounts = [Decimal(value) for value in re.findall(r"<TRNAMT>([^<\r\n]+)", text)]
    assert amounts == [t.amount for t in statement.transactions]
    assert re.findall(r"<DTPOSTED>(\d{8})", text) == ["20260101", "20260101"]
    assert "&lt;merchant&gt; &amp; café" in text
    assert ("<CREDITCARDMSGSRSV1>" if direction == "liability" else "<BANKMSGSRSV1>") in text


def test_ofx_missing_summary_never_guesses():
    from stmtconv.errors import StmtconvError

    statement = monthly(1)
    statement.summary.closing = None
    with pytest.raises(StmtconvError):
        ofx.serialize(statement, "20260101000000")


@pytest.mark.parametrize("direction,name", [("asset", "bank"), ("liability", "card")])
def test_ofx_manual_golden_and_sgml_parse(direction, name):
    import xml.etree.ElementTree as ET
    from pathlib import Path

    statement = monthly(1)
    statement.summary.direction = direction
    output = ofx.serialize(statement, "20260101000000")
    assert output == (Path(__file__).resolve().parent / "golden" / f"{name}.ofx").read_bytes()
    sgml = output.decode("cp1252").split("<OFX>", 1)[1]
    leaves = {
        "CODE",
        "SEVERITY",
        "DTSERVER",
        "LANGUAGE",
        "TRNUID",
        "CURDEF",
        "ACCTID",
        "BANKID",
        "ACCTTYPE",
        "DTSTART",
        "DTEND",
        "TRNTYPE",
        "DTPOSTED",
        "TRNAMT",
        "FITID",
        "NAME",
        "MEMO",
        "BALAMT",
        "DTASOF",
    }
    xml = re.sub(
        r"<([A-Z0-9]+)>([^<]+)",
        lambda m: f"<{m[1]}>{m[2].strip()}</{m[1]}>" if m[1] in leaves else m[0],
        "<OFX>" + sgml,
    )
    parsed = ET.fromstring(xml)
    assert parsed.find(".//CURDEF").text == "USD"
    assert Decimal(parsed.find(".//BALAMT").text) == statement.summary.closing
    assert Decimal(parsed.find(".//STMTTRN/TRNAMT").text) == statement.transactions[0].amount


def test_same_mask_is_not_sufficient_identity():
    first, second = monthly(1), monthly(2)
    first.summary.account_key = None
    second.summary.account_key = None
    assert any(
        i.code == "ACCOUNT_IDENTITY_UNCONFIRMED"
        for i in merge([first, second], Decimal("0.01")).issues
    )
    assert not merge(
        [first, second], Decimal("0.01"), operator_group="owner-confirmed-same-account"
    ).issues
    second.summary.account_key = "b" * 64
    first.summary.account_key = "a" * 64
    assert any(
        i.code == "ACCOUNT_GROUP_UNKNOWN" for i in merge([first, second], Decimal("0.01")).issues
    )
