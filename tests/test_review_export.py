from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from stmtconv.catalog import load_catalog
from stmtconv.config import load_settings
from stmtconv.core.models import Statement, StatementSummary, Transaction
from stmtconv.core.validate import validate
from stmtconv.demo import generate
from stmtconv.errors import StmtconvError
from stmtconv.export import csv_writer, excel
from stmtconv.export.service import export
from stmtconv.extract.service import extract, read_statements
from stmtconv.intake.service import intake
from stmtconv.orders import store
from stmtconv.review import service

REPO = Path(__file__).resolve().parents[1]


def sample():
    summary = StatementSummary(
        opening=Decimal("1000"),
        closing=Decimal("2987.50"),
        total_debits=Decimal("12.50"),
        total_credits=Decimal("2000"),
        period_start=date(2026, 1, 1),
        period_end=date(2026, 1, 31),
        account_mask="****1234",
    )
    transactions = [
        Transaction(
            id="row-a",
            date=date(2026, 1, 2),
            description="Fake café, grocery",
            debit=Decimal("12.50"),
            balance=Decimal("987.50"),
            source_file="fake.pdf",
            page=1,
            engine="text",
            raw_text="fake",
        ),
        Transaction(
            id="row-b",
            date=date(2026, 1, 3),
            description="Fake salary",
            credit=Decimal("2000.00"),
            balance=Decimal("2987.50"),
            source_file="fake.pdf",
            page=1,
            engine="text",
            raw_text="fake",
        ),
    ]
    return validate(
        Statement(id="s0001", summary=summary, transactions=transactions, pages=[1]),
        Decimal("0.01"),
    )


@pytest.mark.parametrize("format_name", ["qb_csv3", "qb_csv4", "xero_csv", "csv"])
def test_csv_matches_manual_golden(format_name, config_dir):
    spec = load_catalog(config_dir).exports.formats[format_name]
    output = csv_writer.serialize(
        sample(), spec, "%Y-%m-%d" if format_name == "csv" else "%m/%d/%Y", bom=format_name == "csv"
    )
    assert output == (REPO / "tests/golden" / f"{format_name}.csv").read_bytes()
    if format_name.startswith("qb_"):
        assert b",0," not in output
        assert b"$" not in output


def test_excel_dates_money_and_mismatch_fill(config_dir, tmp_path):
    statement = sample()
    statement.transactions[0].debit = Decimal("99")
    statement = validate(statement, Decimal("0.01"))
    path = tmp_path / "output.xlsx"
    excel.write(statement, path, load_catalog(config_dir).exports.formats["excel"], True)
    book = load_workbook(path)
    assert book["Transactions"]["A2"].value.date() == date(2026, 1, 2)
    assert book["Transactions"]["C2"].value == 99
    assert book["Transactions"]["C2"].data_type == "n"
    assert book["Transactions"]["A2"].fill.fgColor.rgb.endswith("FFF2CC")
    assert "not fully verified" in str(dict(book["Summary"].values)["verification"])


def setup_order(config_dir, corrupt=False):
    settings = load_settings(config_dir)
    order = store.create(settings.workspace)
    folder = store.root(settings.workspace, order.order_id) / "input"
    path, _ = generate(folder, count=5, corrupt_index=1 if corrupt else None)
    path.with_suffix(".json").unlink()
    intake(settings, order.order_id)
    extract(settings, order.order_id, engine="text")
    return settings, order.order_id


def test_review_fixes_mismatch_and_stale_reapply_rejected(config_dir):
    settings, order_id = setup_order(config_dir, True)
    path = service.write(settings, order_id)[0]
    book = load_workbook(path)
    sheet = book["Review"]
    sheet["O3"] = 2
    sheet["R3"] = "fix"
    book.save(path)
    result = service.apply(settings, order_id)
    assert result[0].verdict == "VERIFIED_WITH_FIXES"
    assert store.load(settings.workspace, order_id).manual_fixes == 1
    assert result[0].transactions[1].fixed_by == "review"
    with pytest.raises(StmtconvError, match="stale"):
        service.apply(settings, order_id)


def test_invalid_review_batch_changes_nothing(config_dir):
    settings, order_id = setup_order(config_dir, True)
    before = read_statements(settings, order_id)[0].model_dump_json()
    path = service.write(settings, order_id)[0]
    book = load_workbook(path)
    book["Review"]["O3"] = 2
    book["Review"]["R3"] = "fix"
    book["Review"]["O4"] = "=1+1"
    book["Review"]["R4"] = "fix"
    book.save(path)
    with pytest.raises(StmtconvError, match="sheet row 4"):
        service.apply(settings, order_id)
    assert read_statements(settings, order_id)[0].model_dump_json() == before
    assert store.load(settings.workspace, order_id).manual_fixes == 0


def test_export_gates_and_explicit_disclosure(config_dir):
    settings, order_id = setup_order(config_dir, True)
    specs = load_catalog(config_dir).exports
    with pytest.raises(StmtconvError, match="flagged"):
        export(settings, order_id, specs)
    with pytest.raises(StmtconvError, match="spot-check"):
        export(settings, order_id, specs, allow_unverified=True)
    paths = export(settings, order_id, specs, allow_unverified=True, skip_spotcheck=True)
    assert len(paths) == 2
    order = store.load(settings.workspace, order_id)
    assert order.exported_unverified and order.spot_check.skipped
    assert order.status == "exported"


def test_passed_spotcheck_and_qb_splitting(config_dir):
    settings, order_id = setup_order(config_dir)
    settings.qb_csv_max_rows = 2
    with store.edit(settings.workspace, order_id) as order:
        order.outputs = ["qb_csv3"]
    service.record_spotcheck(settings, order_id, True, "synthetic source checked")
    paths = export(settings, order_id, load_catalog(config_dir).exports)
    assert len(paths) == 3
    assert all(len(path.read_text().splitlines()) <= 3 for path in paths)


def test_csv_formula_like_description_is_escaped(config_dir):
    statement = sample()
    statement.transactions[0].description = '=HYPERLINK("fake")'
    output = csv_writer.serialize(
        statement, load_catalog(config_dir).exports.formats["qb_csv3"], "%m/%d/%Y"
    )
    assert b"'=HYPERLINK" in output


def test_fixing_balance_does_not_hide_an_unreadable_debit(config_dir):
    from stmtconv.core.models import RawRow
    from stmtconv.core.normalize import normalize
    from stmtconv.extract.service import persist

    settings, order_id = setup_order(config_dir)
    statement = validate(
        normalize(
            "s0001",
            [
                RawRow(
                    date="2026-01-01",
                    description="Fake unknown",
                    debit="broken",
                    balance="1000",
                    source_file="fake.pdf",
                    raw_text="broken",
                )
            ],
            StatementSummary(
                opening=Decimal("1000"),
                closing=Decimal("1000"),
                total_debits=Decimal(0),
                total_credits=Decimal(0),
            ),
        ),
        Decimal("0.01"),
    )
    persist(settings, order_id, [statement])
    with store.edit(settings.workspace, order_id) as order:
        store.transition(order, "needs_review")
    path = service.write(settings, order_id)[0]
    book = load_workbook(path)
    book["Review"]["Q2"] = 1000
    book["Review"]["R2"] = "fix"
    book.save(path)
    result = service.apply(settings, order_id)[0]
    assert result.verdict == "NEEDS_REVIEW"
    assert "AMOUNT_UNPARSABLE" in result.transactions[0].flags


def test_review_allows_negative_balances_but_not_negative_debits():
    assert service.fix_money("-10.50", negative=True) == Decimal("-10.50")
    with pytest.raises(ValueError):
        service.fix_money("-10.50")
