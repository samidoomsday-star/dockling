from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from PIL import Image
from pydantic import ValidationError
from reportlab.pdfgen.canvas import Canvas

from stmtconv.catalog import load_catalog
from stmtconv.config import load_settings
from stmtconv.core.dates import parse_date
from stmtconv.core.models import StatementSummary, Transaction
from stmtconv.demo import generate
from stmtconv.demo_assets import assets
from stmtconv.errors import StmtconvError
from stmtconv.export.csv_writer import cells
from stmtconv.extract.service import extract
from stmtconv.intake.service import intake
from stmtconv.launcher import launch
from stmtconv.orders import ledger, store
from stmtconv.pipeline import run
from stmtconv.privacy.deletion import close
from stmtconv.review.service import record_spotcheck


@pytest.mark.parametrize("year,expected", [(2024, date(2024, 2, 29)), (2023, None)])
def test_missing_year_february_29(year, expected):
    parsed, flags = parse_date("29 Feb", "DMY", date(year, 2, 1), date(year, 3, 1))
    assert parsed == expected
    assert flags == ([] if expected else ["DATE_UNPARSABLE"])


def test_float_money_rejected_and_zero_balance_retained():
    with pytest.raises(ValidationError):
        StatementSummary.model_validate_json('{"opening":0.1}')
    row = Transaction(
        id="zero",
        date=date(2026, 1, 1),
        description="Fake",
        debit=Decimal(0),
        credit=Decimal(0),
        balance=Decimal(0),
        source_file="fake.pdf",
        page=1,
        engine="text",
        raw_text="",
    )
    assert cells(row, "%Y-%m-%d")["Balance"] == "0.00"
    assert cells(row, "%Y-%m-%d")["Debit"] == "0.00"
    assert cells(row, "%Y-%m-%d", blank_zero=True)["Debit"] == ""
    assert cells(row, "%Y-%m-%d", blank_zero=True)["Balance"] == "0.00"


def test_inbox_moves_then_pauses_and_resumes_to_delivery(config_dir, tmp_path):
    settings, catalog = load_settings(config_dir), load_catalog(config_dir)
    inbox = tmp_path / "inbox"
    pdf, _ = generate(inbox, count=5)
    pdf.with_suffix(".json").unlink()
    order_id, message, paths = launch(settings, catalog, inbox)
    assert message == "Run spotcheck, then run again" and paths == []
    assert list(inbox.iterdir()) == []
    assert (store.root(settings.workspace, order_id) / "input" / pdf.name).is_file()
    record_spotcheck(settings, order_id, True)
    message, paths = run(settings, catalog, order_id)
    assert message == "Delivered" and paths[0].is_file()
    close(settings, order_id)


def test_invalid_inbox_keeps_files_and_creates_no_order(config_dir, tmp_path):
    settings = load_settings(config_dir)
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    (inbox / "unknown.txt").write_text("keep")
    with pytest.raises(StmtconvError):
        launch(settings, load_catalog(config_dir), inbox)
    assert (inbox / "unknown.txt").read_text() == "keep"
    assert store.listing(settings.workspace) == []


@pytest.mark.parametrize("status", ["created", "intake_done", "reviewed"])
def test_abandoned_order_explicit_cleanup(config_dir, status):
    settings = load_settings(config_dir)
    order = store.create(settings.workspace, alias="FAKE_PRIVATE_ALIAS")
    root = store.root(settings.workspace, order.order_id)
    if status != "created":
        pdf, _ = generate(root / "input", count=5)
        pdf.with_suffix(".json").unlink()
        intake(settings, order.order_id)
        if status == "reviewed":
            extract(settings, order.order_id, engine="text")
    else:
        (root / "input" / "private.txt").write_text("fake private content")
    with pytest.raises(StmtconvError, match="abandon"):
        close(settings, order.order_id)
    close(settings, order.order_id, abandon=True)
    assert [p.name for p in root.iterdir()] == ["order.json"]
    assert "FAKE_PRIVATE_ALIAS" not in (root / "order.json").read_text()
    assert ledger.read(settings.workspace)[0].outcome == "abandoned"


def test_invoice_header_rejected_but_original_preserved(config_dir):
    settings = load_settings(config_dir)
    order = store.create(settings.workspace)
    path = store.root(settings.workspace, order.order_id) / "input" / "invoice.pdf"
    canvas = Canvas(str(path))
    canvas.drawString(40, 750, "Tax Invoice #FAKE123")
    canvas.drawString(40, 730, "Synthetic purchased services payment due")
    canvas.save()
    original = path.read_bytes()
    with pytest.raises(StmtconvError, match="invoice|Invoice"):
        intake(settings, order.order_id)
    assert path.read_bytes() == original
    assert store.load(settings.workspace, order.order_id).status == "created"


def test_blank_page_is_accounted_for_without_missing_transactions(config_dir):
    import pypdfium2 as pdfium

    settings = load_settings(config_dir)
    order = store.create(settings.workspace)
    folder = store.root(settings.workspace, order.order_id) / "input"
    pdf, _ = generate(folder, count=5)
    pdf.with_suffix(".json").unlink()
    with pdfium.PdfDocument(pdf) as document:
        document.new_page(612, 792)
        document.save(folder / "with-blank.pdf")
    pdf.unlink()
    intake(settings, order.order_id)
    result = extract(settings, order.order_id, engine="text")[0]
    assert result.verdict == "VERIFIED"
    assert 2 in result.summary_pages
    assert len(result.transactions) == 5


def test_demo_assets_are_valid_synthetic_outputs(config_dir, tmp_path):
    folder = tmp_path / "assets"
    paths = assets(load_settings(config_dir), load_catalog(config_dir), folder)
    assert len(paths) == 4 and all(p.is_file() for p in paths)
    workbook = load_workbook(folder / "synthetic-sample.xlsx")
    assert "Summary" in workbook.sheetnames
    for name in ["before.png", "after.png"]:
        with Image.open(folder / name) as image:
            image.verify()
    with pytest.raises(StmtconvError, match="empty folder"):
        assets(load_settings(config_dir), load_catalog(config_dir), folder)


def test_processing_requires_offline_settings(config_dir):
    settings = load_settings(config_dir)
    settings.offline = False
    with pytest.raises(StmtconvError, match="OFFLINE|OFFLINE=true"):
        extract(settings, "unused")
