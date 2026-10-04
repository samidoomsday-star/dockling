"""Client workbooks with numeric money, source provenance and honest summaries."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from stmtconv.catalog import ExportFormat
from stmtconv.core.models import Statement
from stmtconv.export.csv_writer import safe_text
from stmtconv.export.templates import render


def write(
    statement: Statement,
    path: Path,
    spec: ExportFormat,
    unverified: bool = False,
    categorize: bool = False,
    skipped_spotcheck: bool = False,
    date_format: str = "%Y-%m-%d",
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Transactions"
    columns = [
        *spec.columns,
        *(["Category"] if categorize and "Category" not in spec.columns else []),
    ]
    sheet.append(columns)
    for row in statement.transactions:
        values: dict[str, object] = {
            "Date": row.date,
            "Description": safe_text(row.description),
            "Debit": row.debit,
            "Credit": row.credit,
            "Balance": row.balance,
            "Check": row.check.status,
            "Page": row.page,
            "Source file": safe_text(row.source_file),
            "Category": safe_text(row.category or ""),
        }
        sheet.append([values.get(column) for column in columns])
        for cell, column in zip(sheet[sheet.max_row], columns, strict=True):
            if column in {"Debit", "Credit", "Balance"}:
                cell.number_format = "#,##0.00"
            elif column == "Date":
                cell.number_format = (
                    date_format.replace("%Y", "yyyy").replace("%m", "mm").replace("%d", "dd")
                )
            if row.check.status == "MISMATCH":
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = (
            45 if column[0].value == "Description" else 18
        )
    summary = workbook.create_sheet("Summary")
    data: dict[str, object] = {
        "account": statement.summary.account_mask or "Unknown",
        "period_start": statement.summary.period_start,
        "period_end": statement.summary.period_end,
        "currency": statement.summary.currency,
        "opening": statement.summary.opening,
        "closing": statement.summary.closing,
        "total_debits": statement.debit_total,
        "total_credits": statement.credit_total,
        "rows": len(statement.transactions),
        "mismatches": sum(t.check.status == "MISMATCH" for t in statement.transactions),
        "fixes": statement.manual_fixes,
        "verdict": statement.verdict,
        "disclosure": "not fully verified" if unverified else "financial reconciliation passed",
        "spotcheck_skipped": skipped_spotcheck,
    }
    for key, value in data.items():
        summary.append([key, value])
    summary.append(["verification", render("verification_summary.md", data)])
    issues = workbook.create_sheet("Issues")
    issues.append(["Row ID", "Page", "Flag", "Resolution"])
    for flag in statement.flags:
        issues.append(["statement", None, flag, ""])
    for row in statement.transactions:
        for flag in row.flags:
            issues.append([row.id, row.page, flag, row.fixed_by or ""])
    workbook.save(path)
