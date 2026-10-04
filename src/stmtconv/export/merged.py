"""Monthly sheets plus full sequence, summary and explicit continuity issues."""

from pathlib import Path

from openpyxl import Workbook

from stmtconv.core.merge import MergeResult
from stmtconv.export.csv_writer import safe_text


def write(result: MergeResult, path: Path) -> None:
    workbook = Workbook()
    all_rows = workbook.active
    assert all_rows is not None
    all_rows.title = "All Transactions"
    columns = [
        "Date",
        "Description",
        "Debit",
        "Credit",
        "Balance",
        "Category",
        "Statement",
        "Page",
        "Source file",
    ]
    all_rows.append(columns)
    summary = workbook.create_sheet("Summary")
    summary.append(["Statement", "Opening", "Closing", "Debits", "Credits", "Rows", "Verdict"])
    summary.append(["merge_verdict", "NEEDS_REVIEW" if result.issues else "VERIFIED_CONTINUITY"])
    for index, statement in enumerate(result.statements, 1):
        name = (
            statement.summary.period_start.strftime("%Y-%m")
            if statement.summary.period_start
            else f"Period-{index}"
        )
        if name in workbook.sheetnames:
            name += f"-{index}"
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for row in statement.transactions:
            values = [
                row.date,
                safe_text(row.description),
                row.debit,
                row.credit,
                row.balance,
                safe_text(row.category or ""),
                statement.id,
                row.page,
                safe_text(row.source_file),
            ]
            sheet.append(values)
            all_rows.append(values)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        summary.append(
            [
                statement.id,
                statement.summary.opening,
                statement.summary.closing,
                statement.debit_total,
                statement.credit_total,
                len(statement.transactions),
                statement.verdict,
            ]
        )
    issues = workbook.create_sheet("Issues")
    issues.append(["Code", "Statement", "Difference"])
    for issue in result.issues:
        issues.append([issue.code, issue.statement_id, issue.difference])
    all_rows.freeze_panes = "A2"
    all_rows.auto_filter.ref = all_rows.dimensions
    workbook.save(path)
