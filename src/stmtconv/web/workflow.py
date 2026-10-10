"""Offline Phase 3 computation, using the same domain commands and six CLI writers."""

import hashlib
import random
from io import BytesIO
from pathlib import Path
from typing import cast
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from stmtconv.catalog import Categories, Exports, load_catalog
from stmtconv.config import Settings, resolve_config_dir
from stmtconv.core.categorize import Rule, categorize
from stmtconv.core.merge import merge
from stmtconv.core.models import Statement
from stmtconv.core.validate import validate
from stmtconv.errors import StmtconvError
from stmtconv.export import csv_writer, excel, merged, ofx
from stmtconv.extract.service import version
from stmtconv.review.commands import apply_edits
from stmtconv.review.service import HEADERS
from stmtconv.review.service import digest as statement_digest
from stmtconv.review.workbook import read_edits
from stmtconv.web.pipeline import ArtifactBundle, PipelineResult, PipelineTask, digest
from stmtconv.web.schemas import Output


def domains(records: list[dict[str, object]]) -> list[Statement]:
    return [Statement.model_validate(r["domain"]) for r in records]


def sample_ids(job_id: str, statements: list[Statement], size: int = 10) -> list[str]:
    rng = random.Random(job_id)
    ids = []
    for statement in statements:
        sampled = rng.sample(statement.transactions, min(size, len(statement.transactions)))
        for row in [*sampled, *(r for r in statement.transactions if r.fixed_by == "review")]:
            key = statement.id + ":" + row.id
            if key not in ids:
                ids.append(key)
    return ids


def source_basis(statements: list[Statement]) -> str:
    return digest(
        [
            {
                "id": s.id,
                "summary": s.summary.model_dump(mode="json"),
                "rows": [
                    {
                        k: v
                        for k, v in t.model_dump(mode="json").items()
                        if k not in {"category", "check", "source_reviewed"}
                    }
                    for t in s.transactions
                ],
            }
            for s in statements
        ]
    )


def source_valid(
    state: dict[str, object], statements: list[Statement], job_id: str, ai: bool = False
) -> bool:
    required = (
        [s.id + ":" + t.id for s in statements for t in s.transactions if t.engine == "ai"]
        if ai
        else sample_ids(job_id, statements)
    )
    if ai and not required:
        return True
    record = state.get("ai_source" if ai else "spotcheck")
    return (
        isinstance(record, dict)
        and record.get("basis") == source_basis(statements)
        and record.get("row_ids") == required
        and record.get("passed") is True
    )


def export_gate(task: PipelineTask, settings: Settings) -> list[Statement]:
    statements = [validate(s, settings.balance_tolerance) for s in domains(task.statements)]
    date_format = {"ISO": "%Y-%m-%d", "MM/DD/YYYY": "%m/%d/%Y", "DD/MM/YYYY": "%d/%m/%Y"}[
        str(task.options.get("output_date_format", "ISO"))
    ]
    specs = (
        Exports.model_validate(task.runtime_config["exports"])
        if task.runtime_config.get("exports")
        else load_catalog(resolve_config_dir()).exports
    )
    for output in cast(list[Output], task.options["outputs"]):
        if output not in {"csv", "ofx"} and date_format not in specs.formats[output].date_formats:
            raise StmtconvError(
                "EXPORT_DATE_FORMAT",
                "Choose MM/DD/YYYY or DD/MM/YYYY when exporting QuickBooks or Xero files.",
            )
    if "ofx" in cast(list[Output], task.options["outputs"]) and any(
        s.summary.account_key is None and not task.review_state.get("account_group")
        for s in statements
    ):
        raise StmtconvError(
            "OFX_ACCOUNT_ID", "Confirm the account under download options before generating OFX."
        )
    if not statements or any(s.verdict in {"NEEDS_REVIEW", "UNVERIFIABLE"} for s in statements):
        raise StmtconvError(
            "EXPORT_FINANCIAL", "Resolve missing or failed financial checks before exporting."
        )
    if not source_valid(task.review_state, statements, str(task.job_id)) or not source_valid(
        task.review_state, statements, str(task.job_id), True
    ):
        raise StmtconvError("EXPORT_SOURCE", "Complete current source checks before exporting.")
    return statements


def review_book(statement: Statement, binding: list[object]) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Review"
    sheet.append(HEADERS)
    for row in statement.transactions:
        sheet.append(
            [
                row.id,
                row.date,
                csv_writer.safe_text(row.description),
                row.debit,
                row.credit,
                row.balance,
                row.page,
                row.engine,
                row.check.status,
                row.check.expected,
                row.check.difference,
                ";".join(row.flags),
                None,
                None,
                None,
                None,
                None,
                "keep",
                None,
            ]
        )
        if row.check.status == "MISMATCH" or row.flags:
            for cell in sheet[sheet.max_row]:
                cell.fill = PatternFill(
                    "solid", fgColor="FFF2CC" if row.check.status == "MISMATCH" else "FCE4D6"
                )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    metadata = book.create_sheet("_Metadata")
    metadata.append(binding)
    metadata.sheet_state = "hidden"
    issues = book.create_sheet("Issues")
    issues.append(["Flag"])
    for flag in statement.flags:
        issues.append([flag])
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def compute_workflow(
    task: PipelineTask, documents: dict[str, bytes], scratch: Path, settings: Settings
) -> PipelineResult:
    statements = domains(task.statements)
    artifacts: list[ArtifactBundle] = []
    manifest: list[dict[str, object]] = []

    def output(kind: str, name: str, data: bytes, **metadata: object) -> None:
        if (
            not 0 < len(data) <= 8 * 1024**2
            or len(manifest) >= 100
            or sum(len(a.data) for a in artifacts) + len(data) > 20 * 1024**2
        ):
            raise StmtconvError(
                "OUTPUT_LIMIT", "This output exceeds the bounded test worker limit."
            )
        aid = uuid4()
        artifacts.append(ArtifactBundle(aid, kind, data))
        manifest.append(
            {
                "id": str(aid),
                "name": name,
                "kind": kind,
                "sha256": hashlib.sha256(data).hexdigest(),
                "bytes": len(data),
                **metadata,
            }
        )

    if task.action == "review_workbook":
        for statement in statements:
            binding = [
                str(task.workspace_id),
                str(task.job_id),
                task.revision,
                statement.id,
                statement_digest(statement),
            ]
            output(
                "review_workbook",
                statement.id + "-review.xlsx",
                review_book(statement, binding),
                binding=binding,
                statement_id=statement.id,
            )
    elif task.action == "review_apply":
        statement_id = task.payload["statement_id"]
        original = next(s for s in statements if s.id == statement_id)
        edits = read_edits(
            documents["workbook"], original, cast(list[object], task.payload["binding"])
        )
        if not edits:
            raise StmtconvError(
                "REVIEW_EMPTY", "Choose fix, delete or insert_after before applying a workbook."
            )
        changed = apply_edits(statements, edits, settings)
        records = [
            {**r, "domain": s.model_dump(mode="json"), "version": version([s])}
            for r, s in zip(task.statements, changed, strict=True)
        ]
        return PipelineResult(
            {
                "state": "succeeded",
                "digest": digest(records),
                "statements": records,
                "edits": [e.model_dump(mode="json", exclude_unset=True) for e in edits],
            }
        )
    elif task.action == "export":
        statements = export_gate(task, settings)
        rules = Categories.model_validate({"rules": task.payload.get("rules", [])})
        if task.options.get("categorize"):
            statements = [
                categorize(s, [Rule(r.category, r.match, r.direction) for r in rules.rules])
                for s in statements
            ]
        identity = task.review_state.get("account_group")
        if identity:
            for s in statements:
                if s.summary.account_key is None:
                    s.summary.account_key = str(identity)
        specs = (
            Exports.model_validate(task.runtime_config["exports"])
            if task.runtime_config.get("exports")
            else load_catalog(resolve_config_dir()).exports
        )
        date_format = {"ISO": "%Y-%m-%d", "MM/DD/YYYY": "%m/%d/%Y", "DD/MM/YYYY": "%d/%m/%Y"}[
            str(task.options.get("output_date_format", "ISO"))
        ]
        for s in statements:
            for format_name in cast(list[Output], task.options["outputs"]):
                spec = specs.formats[format_name]
                if format_name == "excel":
                    path = scratch / (s.id + ".xlsx")
                    excel.write(
                        s,
                        path,
                        spec,
                        categorize=bool(task.options.get("categorize")),
                        date_format=date_format,
                    )
                    output(
                        "export",
                        path.name,
                        path.read_bytes(),
                        format=format_name,
                        statement_id=s.id,
                    )
                elif format_name == "ofx":
                    output(
                        "export",
                        s.id + ".ofx",
                        ofx.serialize(s, str(task.payload["created_at"])),
                        format=format_name,
                        statement_id=s.id,
                    )
                else:
                    limit = (
                        settings.qb_csv_max_rows
                        if format_name.startswith("qb_")
                        else max(1, len(s.transactions))
                    )
                    for i in range(0, len(s.transactions), limit):
                        name = (
                            s.id
                            + "-"
                            + format_name
                            + ("-part" + str(i // limit + 1) if len(s.transactions) > limit else "")
                            + ".csv"
                        )
                        output(
                            "export",
                            name,
                            csv_writer.serialize(
                                s,
                                spec,
                                "%Y-%m-%d" if format_name == "csv" else date_format,
                                s.transactions[i : i + limit],
                                format_name == "csv",
                                format_name.startswith("qb_"),
                            ),
                            format=format_name,
                            statement_id=s.id,
                        )
        if task.options.get("merge"):
            combination = merge(
                statements, settings.balance_tolerance, str(identity) if identity else None
            )
            if combination.issues:
                raise StmtconvError(
                    "MERGE_BLOCKED", "Confirm one account and resolve all merge continuity issues."
                )
            path = scratch / "monthly-merged.xlsx"
            merged.write(combination, path)
            output("export", path.name, path.read_bytes(), format="merged_excel")
    elif task.action == "delivery":
        export_gate(task, settings)
        exports = cast(list[dict[str, object]], task.payload["exports"])
        buffer = BytesIO()
        with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
            for item in exports:
                content = documents[str(item["id"])]
                if hashlib.sha256(content).hexdigest() != item["sha256"]:
                    raise StmtconvError(
                        "ARTIFACT_INTEGRITY", "Regenerate the output set before delivery."
                    )
                archive.writestr(str(item["name"]), content)
            archive.writestr(
                "VERIFICATION-SUMMARY.txt",
                "Financial checks passed. The required source sample and all fixed rows were compared. Sampling does not prove every date or description. AI-source review, when applicable, covers all AI rows.\n",
            )
            archive.writestr(
                "DELIVERY-NOTE.txt",
                str(
                    cast(dict[str, object], task.runtime_config.get("templates", {})).get(
                        "delivery_note", ""
                    )
                )
                + "\nThese outputs were produced from your current reviewed revision. This package is available for download; it is not proof of email delivery or accounting-software import.\n",
            )
            archive.writestr(
                "DATA-HANDLING.txt",
                "Original documents are excluded from this package. Owner removal revokes live access immediately. A live-removal certificate is issued only after verified cleanup; backup expiry is a separate hosting policy, currently undecided.\n",
            )
        output("delivery", "reviewed-outputs.zip", buffer.getvalue())
    else:
        raise StmtconvError("OPERATION_UNSUPPORTED", "This operation is unavailable.")
    return PipelineResult(
        {
            "state": "succeeded",
            "digest": digest(manifest),
            "manifest": {"revision": task.revision, "items": manifest},
        },
        artifacts,
    )
