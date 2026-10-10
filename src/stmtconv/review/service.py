"""Review edits are validated as a complete batch before any change is saved."""

import hashlib
import random
from pathlib import Path
from time import perf_counter

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from stmtconv.config import Settings
from stmtconv.core.models import Statement, Transaction
from stmtconv.core.validate import validate
from stmtconv.export.csv_writer import safe_text
from stmtconv.extract.service import persist, read_statements, version
from stmtconv.orders import store
from stmtconv.review.commands import apply_edits, fix_date, fix_money  # noqa: F401
from stmtconv.review.workbook import read_edits

HEADERS = [
    "Row ID",
    "Date",
    "Description",
    "Debit",
    "Credit",
    "Balance",
    "Page",
    "Engine",
    "Check",
    "Expected",
    "Difference",
    "Flags",
    "Fix Date",
    "Fix Description",
    "Fix Debit",
    "Fix Credit",
    "Fix Balance",
    "Action",
    "Note",
]


def digest(statement: Statement) -> str:
    return hashlib.sha256(statement.model_dump_json().encode()).hexdigest()


def write(settings: Settings, order_id: str) -> list[Path]:
    order = store.load(settings.workspace, order_id)
    store.require(order, {"needs_review", "reviewed"})
    paths = []
    for statement in read_statements(settings, order_id):
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Review"
        sheet.append(HEADERS)
        for row in statement.transactions:
            sheet.append(
                [
                    row.id,
                    row.date,
                    safe_text(row.description),
                    row.debit,
                    row.credit,
                    row.balance,
                    row.page,
                    row.engine,
                    row.check.status,
                    row.check.expected,
                    row.check.difference,
                    ";".join(row.flags),
                    None,
                    None,
                    None,
                    None,
                    None,
                    "keep",
                    None,
                ]
            )
            if row.check.status == "MISMATCH" or row.flags:
                for cell in sheet[sheet.max_row]:
                    cell.fill = PatternFill(
                        "solid", fgColor="FFF2CC" if row.check.status == "MISMATCH" else "FCE4D6"
                    )
        sheet.freeze_panes = "A2"
        metadata = workbook.create_sheet("_Metadata")
        metadata.append([order_id, statement.id, digest(statement)])
        metadata.sheet_state = "hidden"
        issues = workbook.create_sheet("Issues")
        issues.append(["Flag"])
        for flag in statement.flags:
            issues.append([flag])
        path = store.child(
            store.root(settings.workspace, order_id), f"review/{statement.id}-review.xlsx"
        )
        temporary = path.with_suffix(".tmp.xlsx")
        workbook.save(temporary)
        temporary.replace(path)
        paths.append(path)
    return paths


def apply(settings: Settings, order_id: str, confirm_ai_source: bool = False) -> list[Statement]:
    started = perf_counter()
    with store.edit(settings.workspace, order_id) as order:
        store.require(order, {"needs_review", "reviewed"})
        originals = read_statements(settings, order_id)
        edits = []
        for original in originals:
            path = store.child(
                store.root(settings.workspace, order_id), f"review/{original.id}-review.xlsx"
            )
            edits.extend(
                read_edits(path.read_bytes(), original, [order_id, original.id, digest(original)])
            )
        changed = apply_edits(originals, edits, settings)
        if confirm_ai_source:
            for result in changed:
                for row in result.transactions:
                    if row.engine == "ai":
                        row.source_reviewed = True
                        row.fixed_by = row.fixed_by or "ai"
                        row.flags = [f for f in row.flags if f != "AI_RESULT_UNTRUSTED"]
                result.flags = [f for f in result.flags if f != "AI_RESULT_UNTRUSTED"]
            changed = [validate(result, settings.balance_tolerance) for result in changed]
        persist(settings, order_id, changed)
        order.manual_fixes += len(edits)
        order.timings["review"] = order.timings.get("review", 0) + perf_counter() - started
        order.spot_check = type(order.spot_check)()
        for fact, statement in zip(order.statements, changed, strict=True):
            fact.verdict = statement.verdict
            fact.rows = len(statement.transactions)
            fact.mismatches = sum(t.check.status == "MISMATCH" for t in statement.transactions)
        from stmtconv.orders.models import Status

        target: Status = (
            "needs_review"
            if any(s.verdict in {"NEEDS_REVIEW", "UNVERIFIABLE"} for s in changed)
            else "reviewed"
        )
        if order.status != target:
            store.transition(order, target)
    return changed


def sample(settings: Settings, order_id: str) -> list[Transaction]:
    randomizer = random.Random(order_id)
    rows: list[Transaction] = []
    for statement in read_statements(settings, order_id):
        selected = randomizer.sample(
            statement.transactions, min(settings.spot_check_rows, len(statement.transactions))
        )
        by_id = {
            t.id: t
            for t in [*selected, *(t for t in statement.transactions if t.fixed_by == "review")]
        }
        rows.extend(by_id.values())
    return rows


def record_spotcheck(
    settings: Settings, order_id: str, passed: bool, note: str = ""
) -> list[Transaction]:
    with store.edit(settings.workspace, order_id) as order:
        store.require(order, {"needs_review", "reviewed"})
        rows = sample(settings, order_id)
        order.spot_check.sample_ids = [row.id for row in rows]
        order.spot_check.passed = passed
        order.spot_check.revision = version(read_statements(settings, order_id))
        order.spot_check.note = note
    return rows
