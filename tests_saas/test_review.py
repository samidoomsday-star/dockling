"""Real PG/S3 review, six writers, workbook parity and verified removal failures."""

import base64
import hashlib
from io import BytesIO
from uuid import UUID, uuid4
from zipfile import ZipFile

import pytest
import test_pipeline
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from test_foundation import create
from test_pipeline import run, upload

from stmtconv.config import load_settings
from stmtconv.web.models import Artifact, CanonicalSnapshot, Job
from stmtconv.web.pipeline import compute


@pytest.fixture
def worker(world):
    yield from test_pipeline.worker.__wrapped__(world)


def converted(world, worker, tmp_path):
    c = world["client"]("owner0")
    jid = create(world)["id"]
    assert upload(c, jid).status_code == 200
    assert c.post(f"/api/v1/jobs/{jid}/intake", json={"expected_revision": 2}).status_code == 202
    run(world, worker, tmp_path)
    assert (
        c.post(
            f"/api/v1/jobs/{jid}/operations", json={"expected_revision": 3, "action": "extract"}
        ).status_code
        == 202
    )
    run(world, worker, tmp_path)
    return c, jid


def workflow_run(world, worker, tmp_path, *, scratch=False):
    token = uuid4().hex + uuid4().hex
    task = worker.claim(token)
    assert task

    def call(action, **values):
        return world["app"].state.broker.call(
            {
                "operation": str(task.id),
                "generation": task.generation,
                "token": token,
                "action": action,
                **values,
            }
        )

    if task.action == "close":
        return task, call("close")
    if scratch:
        call("scratch_begin")
    docs = {}
    if task.action == "review_apply":
        docs["workbook"] = base64.b64decode(
            call("read", artifact=task.payload["upload_id"])["data"]
        )
    elif task.action == "delivery":
        docs = {
            i["id"]: base64.b64decode(call("read", artifact=i["id"])["data"])
            for i in task.payload["exports"]
        }
    result = compute(task, docs, {}, tmp_path, load_settings(), 40)
    for a in result.artifacts:
        call("stage", artifact=str(a.id), kind=a.kind, data=base64.b64encode(a.data).decode())
    assert worker.publish(task, token, result.record)
    if scratch:
        call("scratch_end")
    return task, result


def post(c, path, revision, **fields):
    c.headers["Idempotency-Key"] = str(uuid4())
    return c.post(path, json={"expected_revision": revision, **fields})


def source(c, jid, rev):
    path = f"/api/v1/jobs/{jid}/spotcheck"
    sample = c.get(path).json()
    response = post(c, path, rev, row_ids=sample["row_ids"], passed=True)
    assert response.status_code == 200, response.text


def test_browser_corrections_atomic_revision_and_provenance(world, worker, tmp_path):
    c, jid = converted(world, worker, tmp_path)
    row = c.get(f"/api/v1/jobs/{jid}/rows").json()["items"][0]
    path = f"/api/v1/jobs/{jid}/review"
    command = {
        "expected_revision": 4,
        "edits": [
            {
                "statement_id": row["statement_id"],
                "row_id": row["id"],
                "action": "fix",
                "changes": {"description": "Fictional corrected description"},
            }
        ],
    }
    bad = {**command, "edits": [*command["edits"], {**command["edits"][0], "row_id": "foreign"}]}
    assert c.patch(path, json=bad).status_code == 422
    assert c.get(f"/api/v1/jobs/{jid}").json()["revision"] == 4
    for who in ("owner1", "editor1", "admin", "viewer0"):
        assert world["client"](who).patch(path, json=command).status_code == (
            403 if who == "viewer0" else 404
        )
    source(c, jid, 4)
    c.headers["Idempotency-Key"] = str(uuid4())
    saved = c.patch(path, json=command)
    assert saved.status_code == 200, saved.text
    assert saved.json()["revision"] == 5 and not saved.json()["source_checked"]
    assert c.patch(path, json=command).json() == saved.json()
    c.headers["Idempotency-Key"] = str(uuid4())
    assert c.patch(path, json=command).status_code == 409
    row2 = c.get(f"/api/v1/jobs/{jid}/rows").json()["items"][0]
    assert row2["description"] == "Fictional corrected description" and row2["fixed_by"] == "review"
    assert (row2["page"], row2["file_id"], row2["balance"]) == (
        row["page"],
        row["file_id"],
        row["balance"],
    )


def test_workbook_roundtrip_same_commands_and_stale_denial(world, worker, tmp_path):
    c, jid = converted(world, worker, tmp_path)
    path = f"/api/v1/jobs/{jid}/review-workbooks"
    assert post(c, path, 4).status_code == 202
    workflow_run(world, worker, tmp_path, scratch=True)
    item = c.get(path).json()["items"][0]
    raw = c.get(f"/api/v1/jobs/{jid}/artifacts/{item['id']}/download").content
    book = load_workbook(BytesIO(raw))
    book["Review"]["N2"] = "Edited through Excel"
    book["Review"]["R2"] = "fix"
    output = BytesIO()
    book.save(output)
    apply = f"{path}/{item['id']}/apply"
    c.headers["Idempotency-Key"] = str(uuid4())
    response = c.post(
        apply,
        content=output.getvalue(),
        headers={"Content-Type": "application/octet-stream", "X-Expected-Revision": "4"},
    )
    assert response.status_code == 202, response.text
    workflow_run(world, worker, tmp_path)
    assert c.get(f"/api/v1/jobs/{jid}").json()["revision"] == 5
    assert (
        c.get(f"/api/v1/jobs/{jid}/rows").json()["items"][0]["description"]
        == "Edited through Excel"
    )
    assert c.get(f"/api/v1/jobs/{jid}/artifacts/{item['id']}/download").status_code == 404
    c.headers["Idempotency-Key"] = str(uuid4())
    assert (
        c.post(apply, content=output.getvalue(), headers={"X-Expected-Revision": "5"}).status_code
        == 409
    )


def six_outputs(world, worker, tmp_path):
    c, jid = converted(world, worker, tmp_path)
    options = {
        "expected_revision": 4,
        "outputs": ["excel", "csv", "qb_csv3", "qb_csv4", "xero_csv", "ofx"],
        "output_date_format": "MM/DD/YYYY",
        "merge": False,
        "categorize": False,
    }
    result = c.patch(f"/api/v1/jobs/{jid}/output-options", json=options)
    assert result.status_code == 200, result.text
    assert post(c, f"/api/v1/jobs/{jid}/exports", 5).status_code == 409
    assert (
        post(
            c,
            f"/api/v1/jobs/{jid}/account-group",
            5,
            confirmed=True,
            private_group="Fictional checking account",
        ).status_code
        == 200
    )
    source(c, jid, 6)
    response = post(c, f"/api/v1/jobs/{jid}/exports", 6)
    assert response.status_code == 202, response.text
    workflow_run(world, worker, tmp_path, scratch=True)
    return c, jid


def test_all_six_actual_outputs_hashes_package_and_removal(world, worker, tmp_path):
    c, jid = six_outputs(world, worker, tmp_path)
    manifest = c.get(f"/api/v1/jobs/{jid}/exports").json()
    assert c.get(f"/api/v1/jobs/{jid}/statements").json()["items"][0]["source_checked"]
    assert len(manifest["items"]) == 6
    assert {i["format"] for i in manifest["items"]} == {
        "excel",
        "csv",
        "qb_csv3",
        "qb_csv4",
        "xero_csv",
        "ofx",
    }
    downloads = {}
    for item in manifest["items"]:
        path = f"/api/v1/jobs/{jid}/artifacts/{item['id']}/download"
        raw = c.get(path).content
        assert hashlib.sha256(raw).hexdigest() == item["sha256"]
        assert world["client"]("owner1").get(path).status_code == 404
        downloads[item["name"]] = raw
        if item["format"] == "ofx":
            assert b"<OFX>" in raw and b"<FITID>" in raw
        elif item["format"] == "excel":
            book = load_workbook(BytesIO(raw))
            assert book["Transactions"]["A2"].value.year == 2026
            assert isinstance(book["Transactions"]["D2"].value, (int, float))
        else:
            assert b"20.00" in raw and b"2.00" in raw
    assert post(c, f"/api/v1/jobs/{jid}/delivery", 6).status_code == 202
    workflow_run(world, worker, tmp_path)
    package = c.get(f"/api/v1/jobs/{jid}/delivery").json()["items"][0]
    raw = c.get(f"/api/v1/jobs/{jid}/artifacts/{package['id']}/download").content
    with ZipFile(BytesIO(raw)) as archive:
        assert set(archive.namelist()) == {
            *downloads,
            "VERIFICATION-SUMMARY.txt",
            "DELIVERY-NOTE.txt",
            "DATA-HANDLING.txt",
        }
        for name, content in downloads.items():
            assert archive.read(name) == content
    assert post(c, f"/api/v1/jobs/{jid}/close", 6, confirm_removal=True).status_code == 202
    assert c.get(f"/api/v1/jobs/{jid}/rows").status_code == 404
    assert c.get(f"/api/v1/jobs/{jid}/artifacts/{package['id']}/download").status_code == 404
    _, removed = workflow_run(world, worker, tmp_path)
    assert removed["removed"] is True
    privacy = c.get(f"/api/v1/jobs/{jid}/privacy").json()
    assert (
        privacy["deletion_state"] == "removed" and privacy["certificate"]["worker_scratch_verified"]
    )
    assert not privacy["certificate"]["backup_expiry_verified"]
    with Session(world["engine"]) as db:
        assert db.scalar(select(Artifact.id).where(Artifact.job_id == UUID(jid))) is None
        assert (
            db.scalar(select(CanonicalSnapshot.id).where(CanonicalSnapshot.job_id == UUID(jid)))
            is None
        )
        assert db.get(Job, UUID(jid)).options == {}


def test_partial_delete_withholds_certificate_and_retry_verifies(
    world, worker, tmp_path, monkeypatch
):
    c, jid = converted(world, worker, tmp_path)
    response = post(c, f"/api/v1/jobs/{jid}/close", 4, confirm_removal=True)
    assert response.status_code == 409
    response = post(
        c, f"/api/v1/jobs/{jid}/close", 4, confirm_removal=True, abandon_unfinished=True
    )
    assert response.status_code == 202
    store = world["app"].state.broker.store
    original = store.client.delete_object

    def fail(**kwargs):
        raise OSError("Synthetic deletion failure")

    monkeypatch.setattr(store.client, "delete_object", fail)
    _, removed = workflow_run(world, worker, tmp_path)
    assert removed["removed"] is False
    privacy = c.get(f"/api/v1/jobs/{jid}/privacy").json()
    assert privacy["deletion_state"] == "partial" and privacy["certificate"] is None
    monkeypatch.setattr(store.client, "delete_object", original)
    assert (
        post(
            c, f"/api/v1/jobs/{jid}/close", 5, confirm_removal=True, abandon_unfinished=True
        ).status_code
        == 202
    )
    assert workflow_run(world, worker, tmp_path)[1]["removed"]


def test_scratch_receipt_blocks_close_until_cleanup(world, worker, tmp_path):
    c, jid = converted(world, worker, tmp_path)
    assert post(c, f"/api/v1/jobs/{jid}/review-workbooks", 4).status_code == 202
    token = uuid4().hex + uuid4().hex
    task = worker.claim(token)
    broker = world["app"].state.broker
    payload = {"operation": str(task.id), "generation": task.generation, "token": token}
    broker.call({**payload, "action": "scratch_begin"})
    assert (
        post(
            c, f"/api/v1/jobs/{jid}/close", 4, confirm_removal=True, abandon_unfinished=True
        ).status_code
        == 202
    )
    assert worker.claim(uuid4().hex + uuid4().hex) is None
    assert c.get(f"/api/v1/jobs/{jid}/privacy").json()["certificate"] is None
    broker.call({**payload, "action": "scratch_end"})
    assert workflow_run(world, worker, tmp_path)[1]["removed"]


def test_output_edits_keep_source_but_row_edits_revoke_downloads(world, worker, tmp_path):
    c, jid = six_outputs(world, worker, tmp_path)
    item = c.get(f"/api/v1/jobs/{jid}/exports").json()["items"][0]
    assert (
        c.patch(
            f"/api/v1/jobs/{jid}/output-options",
            json={
                "expected_revision": 6,
                "outputs": ["csv"],
                "output_date_format": "ISO",
                "merge": False,
                "categorize": False,
            },
        ).status_code
        == 200
    )
    assert c.get(f"/api/v1/jobs/{jid}/spotcheck").json()["passed"] is True
    assert c.get(f"/api/v1/jobs/{jid}/exports").json()["items"] == []
    assert c.get(f"/api/v1/jobs/{jid}/artifacts/{item['id']}/download").status_code == 404
    row = c.get(f"/api/v1/jobs/{jid}/rows").json()["items"][0]
    c.headers["Idempotency-Key"] = str(uuid4())
    response = c.patch(
        f"/api/v1/jobs/{jid}/review",
        json={
            "expected_revision": 7,
            "edits": [
                {
                    "row_id": row["id"],
                    "statement_id": row["statement_id"],
                    "action": "fix",
                    "changes": {"credit": None},
                }
            ],
        },
    )
    assert response.status_code == 200
    assert c.get(f"/api/v1/jobs/{jid}/spotcheck").json()["passed"] is False
    assert post(c, f"/api/v1/jobs/{jid}/exports", 8).status_code == 409
    assert (
        c.get(f"/api/v1/jobs/{jid}/review-history").json()["items"][0]["command"]["edits"][0][
            "changes"
        ]["credit"]
        is None
    )


def test_category_validation_preview_and_workspace_invalidation(world, worker, tmp_path):
    c, jid = converted(world, worker, tmp_path)
    invalid = {
        "expected_revision": 4,
        "rules": [{"category": "Unsafe", "match": ["regex:(a+)+$"], "direction": "any"}],
    }
    assert c.put(f"/api/v1/jobs/{jid}/categories", json=invalid).status_code == 422
    rules = [{"category": "Fictional spending", "match": ["purchase"], "direction": "debit"}]
    response = post(c, f"/api/v1/jobs/{jid}/categories/preview", 4, rules=rules)
    assert response.status_code == 200, response.text
    assert response.json()["items"][0]["category"] == "Uncategorized"
    assert response.json()["items"][1]["category"] == "Fictional spending"
    c.headers["Idempotency-Key"] = str(uuid4())
    assert (
        c.patch(
            f"/api/v1/jobs/{jid}/output-options",
            json={
                "expected_revision": 4,
                "outputs": ["excel"],
                "output_date_format": "ISO",
                "merge": False,
                "categorize": True,
            },
        ).status_code
        == 200
    )
    source(c, jid, 5)
    c.headers["Idempotency-Key"] = str(uuid4())
    assert (
        c.put("/api/v1/categories", json={"expected_revision": 1, "rules": rules}).status_code
        == 200
    )
    assert c.get(f"/api/v1/jobs/{jid}").json()["revision"] == 6
    assert c.get(f"/api/v1/jobs/{jid}/spotcheck").json()["passed"]
    assert world["client"]("owner1").get("/api/v1/categories").json()["rules"] == []


def test_tampered_export_cannot_download_or_be_packaged(world, worker, tmp_path):
    c, jid = six_outputs(world, worker, tmp_path)
    item = c.get(f"/api/v1/jobs/{jid}/exports").json()["items"][0]
    with Session(world["engine"]) as db:
        artifact = db.get(Artifact, UUID(item["id"]))
        store = world["app"].state.broker.store
        store.client.put_object(
            Bucket=store.bucket, Key=artifact.object_key, Body=b"tampered synthetic bytes"
        )
    assert c.get(f"/api/v1/jobs/{jid}/artifacts/{item['id']}/download").status_code == 503
    assert post(c, f"/api/v1/jobs/{jid}/delivery", 6).status_code == 503
    assert c.get(f"/api/v1/jobs/{jid}/delivery").json()["items"] == []
