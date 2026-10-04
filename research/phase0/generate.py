"""Synthetic research inputs only; not the product's Phase 2 generator."""

import json
import random
from decimal import Decimal
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from openpyxl import Workbook
from PIL import ImageEnhance
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

HEADERS = ["Date", "Description", "Debit", "Credit", "Balance"]


def rows(count=8):
    balance = Decimal("1000.00")
    result = []
    for i in range(count):
        debit = Decimal("11.25") + Decimal(i) if i % 3 else Decimal(0)
        credit = Decimal("50.00") if i % 3 == 0 else Decimal(0)
        balance += credit - debit
        result.append(
            [
                f"{i + 1:02d} Sep 2026",
                f"Synthetic shop {i + 1:02d}" + (" wrapped description" if i == 2 else ""),
                f"{debit:.2f}" if debit else "",
                f"{credit:.2f}" if credit else "",
                f"{balance:.2f}",
            ]
        )
    return result


def make_pdf(path, data, ruled=True, page_size=8):
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle("Synthetic statement — no real client data")
    x = [40, 145, 350, 415, 485]
    for first in range(0, len(data), page_size):
        c.setFont("Helvetica-Bold", 14)
        c.drawString(40, 800, "SYNTHETIC BANK — TEST DATA ONLY")
        c.setFont("Helvetica", 10)
        c.drawString(40, 780, "Account ****1234 | 01 Sep 2026 - 30 Sep 2026")
        c.drawString(40, 764, "Opening balance 1000.00")
        c.setFont("Helvetica-Bold", 10)
        for pos, text in zip(x, HEADERS, strict=True):
            c.drawString(pos, 730, text)
        c.setFont("Helvetica", 10)
        for n, row in enumerate(data[first : first + page_size]):
            y = 695 - n * 38
            for col, (pos, text) in enumerate(zip(x, row, strict=True)):
                if col == 1 and " wrapped" in text:
                    short, tail = text.split(" wrapped", 1)
                    c.drawString(pos, y, short)
                    c.drawString(pos, y - 12, "wrapped" + tail)
                else:
                    c.drawString(pos, y, text)
            if ruled:
                c.line(35, y - 20, 560, y - 20)
        if ruled:
            c.line(35, 718, 560, 718)
        c.drawString(40, 250, f"Closing balance {data[-1][-1]}")
        c.drawString(40, 230, f"Page {first // page_size + 1}")
        c.showPage()
    c.save()


def generate(out):
    out.mkdir(parents=True, exist_ok=True)
    data = rows()
    make_pdf(out / "digital.pdf", data)
    make_pdf(out / "borderless.pdf", data, ruled=False)
    make_pdf(out / "multipage.pdf", rows(12), ruled=False, page_size=6)
    pdf = pdfium.PdfDocument(out / "digital.pdf")
    page = pdf[0]
    image = page.render(scale=200 / 72).to_pil().convert("RGB")
    image.save(out / "scan.png")
    image.save(out / "scanned.pdf", "PDF", resolution=200)
    photo = ImageEnhance.Contrast(image).enhance(0.85).rotate(2, expand=True, fillcolor="white")
    rng = random.Random(2026)
    for _ in range(8000):
        pos = (rng.randrange(photo.width), rng.randrange(photo.height))
        tone = rng.randrange(180, 255)
        photo.putpixel(pos, (tone, tone, tone))
    photo.save(out / "photo.jpg", quality=90)
    page.close()
    pdf.close()
    doc = Document()
    doc.add_heading("Synthetic statement — test data only", 0)
    table = doc.add_table(rows=1, cols=len(HEADERS))
    for cell, text in zip(table.rows[0].cells, HEADERS, strict=True):
        cell.text = text
    for row in data:
        for cell, text in zip(table.add_row().cells, row, strict=True):
            cell.text = text
    doc.save(out / "statement.docx")
    book = Workbook()
    book.active.append(HEADERS)
    for row in data:
        book.active.append(row)
    book.save(out / "statement.xlsx")
    body = "".join(
        "<tr>" + "".join(f"<td>{v}</td>" for v in row) + "</tr>" for row in [HEADERS, *data]
    )
    (out / "statement.html").write_text(
        "<html><body><h1>Synthetic statement</h1><table>" + body + "</table></body></html>"
    )
    truth = {name: data for name in ["digital", "borderless", "scanned", "photo", "statement"]}
    truth["multipage"] = rows(12)
    (out / "truth.json").write_text(json.dumps(truth, indent=2))
    return truth


if __name__ == "__main__":
    generate(Path("explore/inputs"))
