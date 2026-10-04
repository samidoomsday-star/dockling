# Offline/setup flags must precede third-party imports.
# ruff: noqa: E402
"""Offline exploration; output is evidence, not an application readiness claim."""

import importlib.metadata
import json
import logging
import math
import os
import platform
import shutil
import socket
import subprocess
import threading
import time
import uuid
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
os.environ["HF_HUB_OFFLINE"] = "1"
os.environ["TRANSFORMERS_OFFLINE"] = "1"
os.environ["HF_HOME"] = str(ROOT / "explore/hf-cache")
os.environ["OMP_NUM_THREADS"] = "4"

import pdfplumber
import psutil
from docling.datamodel.accelerator_options import (
    AcceleratorDevice,
    AcceleratorOptions,
)
from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import (
    NativePdfPipelineOptions,
    PdfPipelineOptions,
    RapidOcrOptions,
    TableFormerMode,
)
from docling.document_converter import (
    DocumentConverter,
    ImageFormatOption,
    PdfFormatOption,
)
from docling.pipeline.native_pdf_pipeline import NativePdfPipeline
from generate import generate
from openpyxl import Workbook, load_workbook

logging.basicConfig(level=logging.ERROR)


def block_connection(*args, **kwargs):
    raise RuntimeError("Network forbidden during this offline experiment")


def normalized(value):
    return " ".join(str(value).split())


def measured(fn):
    process = psutil.Process()
    peak = [process.memory_info().rss]
    stop = threading.Event()

    def sample():
        while not stop.wait(0.02):
            peak[0] = max(peak[0], process.memory_info().rss)

    thread = threading.Thread(target=sample)
    thread.start()
    start = time.perf_counter()
    try:
        result = fn()
    except Exception as exc:  # noqa: BLE001 — record upstream failure, never mark it passed.
        result = {
            "status": "failed",
            "error_type": type(exc).__name__,
            "detail": str(exc)[:500],
        }
    finally:
        stop.set()
        thread.join()
    result["seconds"] = round(time.perf_counter() - start, 3)
    result["observed_process_peak_mib"] = round(peak[0] / 1024**2, 1)
    return result


def convert(converter, path, truth, out):
    result = converter.convert(path)
    doc = result.document
    folder = out / path.name
    folder.mkdir(parents=True, exist_ok=True)
    doc.save_as_markdown(folder / "document.md")
    doc.save_as_json(folder / "document.json")
    doc.save_as_html(folder / "document.html")
    grids = []
    book = Workbook()
    book.remove(book.active)
    for i, table in enumerate(doc.tables):
        grid = [
            [normalized(cell) for cell in row]
            for row in table.export_to_dataframe(doc=doc).fillna("").values.tolist()
        ]
        grids.extend(grid)
        sheet = book.create_sheet(f"Table {i + 1}")
        for row in grid:
            sheet.append(row)
    excel_rows = 0
    if grids:
        book.save(folder / "tables.xlsx")
        loaded = load_workbook(folder / "tables.xlsx")
        excel_rows = sum(sheet.max_row for sheet in loaded)
        loaded.close()
    wanted = [[normalized(cell) for cell in row] for row in truth]
    matches = sum(row in grids for row in wanted)
    (folder / "confidence.json").write_text(result.confidence.model_dump_json(indent=2))
    return {
        "status": result.status.value,
        "tables": len(doc.tables),
        "table_data_rows": len(grids),
        "exact_truth_rows": matches,
        "expected_rows": len(wanted),
        "all_truth_rows_exact": matches == len(wanted),
        "xlsx_reloaded_rows": excel_rows,
        "exports": ["markdown", "json", "html"],
        "confidence": result.confidence.model_dump(mode="json"),
        "errors": [str(e.error_message)[:240] for e in result.errors],
    }


def baseline(path, truth):
    with pdfplumber.open(path) as doc:
        text = "\n".join(page.extract_text() or "" for page in doc.pages)
        count = len(doc.pages)
    tokens = [row[0] for row in truth]
    return {
        "status": "passed" if all(token in text for token in tokens) else "failed",
        "pages": count,
        "expected_date_tokens": len(tokens),
        "found_date_tokens": sum(token in text for token in tokens),
        "note": "Text accessibility only; no transaction parsing or financial validation.",
    }


def ocr(path, truth, out):
    process = subprocess.run(
        ["tesseract", str(path), "stdout", "-l", "eng", "--psm", "6"],
        check=False,
        capture_output=True,
        text=True,
        timeout=90,
        env={**os.environ, "OMP_THREAD_LIMIT": "4"},
    )
    (out / (path.stem + "-tesseract.txt")).write_text(process.stdout)
    tokens = [row[0] for row in truth]
    return {
        "status": "passed" if process.returncode == 0 else "failed",
        "exit_code": process.returncode,
        "found_date_tokens": sum(token in process.stdout for token in tokens),
        "expected_date_tokens": len(tokens),
        "note": "Standalone alternate OCR, not Docling table conversion or row accuracy.",
    }


def json_safe(value):
    if isinstance(value, float) and not math.isfinite(value):
        return None  # Upstream omits confidence for model-free formats.
    if isinstance(value, dict):
        return {key: json_safe(item) for key, item in value.items()}
    if isinstance(value, list):
        return [json_safe(item) for item in value]
    return value


def main():
    run_id = uuid.uuid4().hex
    out = ROOT / "explore/output" / run_id
    out.mkdir(parents=True)
    inputs = ROOT / "explore/inputs"
    truth = generate(inputs)
    socket.create_connection = block_connection
    socket.socket.connect = block_connection
    socket.socket.connect_ex = block_connection
    tests = {}
    for name in ["digital", "borderless", "multipage"]:
        tests[f"pdfplumber_{name}"] = measured(
            lambda name=name: baseline(inputs / f"{name}.pdf", truth[name])
        )
    office = DocumentConverter(
        allowed_formats=[InputFormat.DOCX, InputFormat.XLSX, InputFormat.HTML]
    )
    for ext in ["docx", "xlsx", "html"]:
        tests[f"docling_{ext}"] = measured(
            lambda ext=ext: convert(office, inputs / f"statement.{ext}", truth["statement"], out)
        )
    native = DocumentConverter(
        allowed_formats=[InputFormat.PDF],
        format_options={
            InputFormat.PDF: PdfFormatOption(
                pipeline_cls=NativePdfPipeline,
                pipeline_options=NativePdfPipelineOptions(),
            )
        },
    )
    tests["docling_native_pdf"] = measured(
        lambda: convert(native, inputs / "digital.pdf", truth["digital"], out / "native")
    )
    if shutil.which("tesseract"):
        for name, filename in [("scanned", "scan.png"), ("photo", "photo.jpg")]:
            tests[f"tesseract_{name}"] = measured(
                lambda name=name, filename=filename: ocr(inputs / filename, truth[name], out)
            )
    models = ROOT / "models"
    options = PdfPipelineOptions(
        artifacts_path=models,
        do_ocr=False,
        do_table_structure=True,
        accelerator_options=AcceleratorOptions(num_threads=4, device=AcceleratorDevice.CPU),
    )
    options.table_structure_options.mode = TableFormerMode.ACCURATE
    ml = DocumentConverter(
        allowed_formats=[InputFormat.PDF],
        format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)},
    )
    tests["docling_ml_digital"] = measured(
        lambda: convert(ml, inputs / "digital.pdf", truth["digital"], out / "ml")
    )
    digital_ok = tests["docling_ml_digital"]["status"] == "success"
    if digital_ok:
        scan_options = options.model_copy(deep=True)
        scan_options.do_ocr = True
        scan_options.ocr_options = RapidOcrOptions(
            backend="torch", lang=["iso:en"], force_full_page_ocr=True
        )
        scanned = DocumentConverter(
            allowed_formats=[InputFormat.PDF, InputFormat.IMAGE],
            format_options={
                InputFormat.PDF: PdfFormatOption(pipeline_options=scan_options),
                InputFormat.IMAGE: ImageFormatOption(pipeline_options=scan_options),
            },
        )
        for name, filename in [("scanned", "scanned.pdf"), ("photo", "photo.jpg")]:
            tests[f"docling_ml_{name}"] = measured(
                lambda name=name, filename=filename: convert(
                    scanned, inputs / filename, truth[name], out / "ml"
                )
            )
        for name in ["borderless", "multipage"]:
            tests[f"docling_ml_{name}"] = measured(
                lambda name=name: convert(ml, inputs / f"{name}.pdf", truth[name], out / "ml")
            )
    else:
        for name in ["scanned", "photo", "borderless", "multipage"]:
            tests[f"docling_ml_{name}"] = {
                "status": "unrun",
                "reason": "Required ML pipeline prerequisite failed; see digital attempt.",
            }
    needed = ["digital", "scanned", "photo"]
    complete = all(tests[f"docling_ml_{name}"]["status"] == "success" for name in needed)
    evidence = {
        "run_id": run_id,
        "platform": platform.platform(),
        "python": platform.python_version(),
        "threads": 4,
        "network_blocked_in_python": True,
        "versions": {
            p: importlib.metadata.version(p)
            for p in [
                "docling",
                "torch",
                "torchvision",
                "rapidocr",
                "pdfplumber",
                "reportlab",
            ]
        },
        "phase0_required_conversions_complete": complete,
        "tests": tests,
    }
    (out / "results.json").write_text(json.dumps(json_safe(evidence), indent=2, allow_nan=False))
    (ROOT / "explore/latest-run.txt").write_text(str(out))
    for name, result in tests.items():
        print(
            name,
            result["status"],
            result.get("seconds", ""),
            "exact rows",
            result.get("exact_truth_rows", "n/a"),
        )
    print("Current evidence:", out / "results.json")
    return 0 if complete else 1


if __name__ == "__main__":
    raise SystemExit(main())
